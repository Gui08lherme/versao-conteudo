import os
from collections import defaultdict
from datetime import datetime, timezone

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ORGANIZATION = os.environ["AZDO_ORGANIZATION"]
PROJECT = os.environ["AZDO_PROJECT"]
QUERY_ID = os.environ["AZDO_QUERY_ID"]
OUTPUT_FILE = os.environ.get("OUTPUT_FILE", "LISTA-CANDIDATE.xlsx")
PAT = os.environ["AZDO_PAT"]
BASE_URL = f"https://dev.azure.com/{ORGANIZATION}/{PROJECT}/_apis"

session = requests.Session()
session.auth = ("", PAT)
session.headers.update({"Accept": "application/json"})


def fail(message):
    raise SystemExit(f"ERRO: {message}")


def check(response, operation):
    if response.ok:
        return
    print(f"Falha em: {operation}")
    print(f"Status HTTP: {response.status_code}")
    try:
        print(response.json().get("message", response.json()))
    except Exception:
        print(response.text[:2000])
    response.raise_for_status()


print("Executando a query Tree...")
query_url = f"{BASE_URL}/wit/wiql/{QUERY_ID}?api-version=7.1"
response = session.get(query_url)
check(response, "executar a query")
query_result = response.json()
relations = query_result.get("workItemRelations", [])
if not relations:
    fail("A query não retornou relações de árvore.")

columns_from_query = query_result.get("columns", [])
refs = {
    str(c.get("name", "")).strip().lower(): str(c.get("referenceName", "")).strip()
    for c in columns_from_query
    if c.get("name") and c.get("referenceName")
}


def ref(names, fallback=None):
    for name in names:
        if refs.get(name.lower()):
            return refs[name.lower()]
    return fallback


fields = {
    "ID": "System.Id",
    "Title": "System.Title",
    "State": "System.State",
    "Priority": ref(["Priority"], "Microsoft.VSTS.Common.Priority"),
    "ExternalTicketIDs": ref([
        "ExternalTicketIDs", "External Ticket IDs",
        "ExternalTicketID", "External Ticket ID"
    ]),
    "Effort": ref(["Effort"], "Microsoft.VSTS.Scheduling.Effort"),
    "Story Points": ref(["Story Points"], "Microsoft.VSTS.Scheduling.StoryPoints"),
    "Assigned To": "System.AssignedTo",
    "Work Item Type": "System.WorkItemType",
    "Tags": "System.Tags",
    "Area Path": "System.AreaPath",
    "Iteration Path": "System.IterationPath",
}

ordered_ids = []
seen = set()
parent_by_id = {}
for relation in relations:
    source = relation.get("source")
    target = relation.get("target")
    if not target or target.get("id") is None:
        continue
    target_id = int(target["id"])
    if target_id not in seen:
        seen.add(target_id)
        ordered_ids.append(target_id)
    if source and source.get("id") is not None:
        parent_by_id[target_id] = int(source["id"])

print(f"Work Items únicos localizados: {len(ordered_ids)}")
if not ordered_ids:
    fail("Nenhum Work Item foi localizado.")


def chunks(values, size=200):
    for start in range(0, len(values), size):
        yield values[start:start + size]


requested_fields = list(dict.fromkeys(v for v in fields.values() if v))
items_by_id = {}
for number, batch in enumerate(chunks(ordered_ids), start=1):
    print(f"Buscando lote {number}...")
    batch_response = session.post(
        f"{BASE_URL}/wit/workitemsbatch?api-version=7.1",
        json={"ids": batch, "fields": requested_fields, "errorPolicy": "Omit"},
    )
    check(batch_response, f"buscar lote {number}")
    for item in batch_response.json().get("value", []):
        items_by_id[int(item["id"])] = item


def value(item, reference):
    if not reference:
        return ""
    result = item.get("fields", {}).get(reference, "")
    if isinstance(result, dict):
        return result.get("displayName") or result.get("uniqueName") or str(result)
    return "" if result is None else result


depth_cache = {}
def depth(item_id, visiting=None):
    if item_id in depth_cache:
        return depth_cache[item_id]
    visiting = set() if visiting is None else visiting
    if item_id in visiting:
        return 0
    visiting.add(item_id)
    parent = parent_by_id.get(item_id)
    result = 0 if parent is None else depth(parent, visiting) + 1
    visiting.remove(item_id)
    depth_cache[item_id] = result
    return result

headers = [
    "ID", "Title 1", "Title 2", "Title 3", "Title 4", "Title 5",
    "State", "Priority", "ExternalTicketIDs", "Effort", "Story Points",
    "Assigned To", "Work Item Type", "Tags", "Area Path",
    "Iteration Path", "Parent"
]
rows = []
for item_id in ordered_ids:
    item = items_by_id.get(item_id)
    if not item:
        continue
    titles = ["", "", "", "", ""]
    titles[min(depth(item_id), 4)] = value(item, fields["Title"])
    rows.append([
        item_id, *titles,
        value(item, fields["State"]),
        value(item, fields["Priority"]),
        value(item, fields["ExternalTicketIDs"]),
        value(item, fields["Effort"]),
        value(item, fields["Story Points"]),
        value(item, fields["Assigned To"]),
        value(item, fields["Work Item Type"]),
        value(item, fields["Tags"]),
        value(item, fields["Area Path"]),
        value(item, fields["Iteration Path"]),
        parent_by_id.get(item_id, ""),
    ])

workbook = Workbook()
worksheet = workbook.active
worksheet.title = "Wonder Squad - Epics (3)"
worksheet.append(headers)
for row in rows:
    worksheet.append(row)

header_fill = PatternFill(fill_type="solid", fgColor="4B1F7A")
header_font = Font(color="FFFFFF", bold=True)
for cell in worksheet[1]:
    cell.fill = header_fill
    cell.font = header_font

worksheet.freeze_panes = "A2"
worksheet.auto_filter.ref = f"A1:Q{worksheet.max_row}"
widths = [12, 42, 42, 42, 42, 42, 16, 12, 25, 12, 14, 38, 20, 38, 38, 42, 14]
for index, width in enumerate(widths, start=1):
    worksheet.column_dimensions[chr(64 + index)].width = width
for row in worksheet.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

workbook.properties.creator = "GitHub Actions"
workbook.properties.description = "Gerado da query Wonder Squad - Epics."
workbook.save(OUTPUT_FILE)
print(f"Arquivo gerado: {OUTPUT_FILE}")
print(f"Linhas de dados: {len(rows)}")
print("Gerado em UTC: " + datetime.now(timezone.utc).isoformat())
